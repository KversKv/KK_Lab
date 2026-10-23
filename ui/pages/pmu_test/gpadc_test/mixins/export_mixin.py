"""GPADCTestUI Excel 导出 Mixin。

从 ui/pages/pmu_test/gpadc_test_ui.py 平移而来，行为零变更。
类装配见 gpadc_test_ui.py。"""

import datetime
from pathlib import Path


class _ExportMixin:
    """Excel 导出 Mixin"""

    def export_result(self):
        import openpyxl
        from openpyxl.styles import (Font, PatternFill, Alignment,
                                     Border, Side)
        from openpyxl.utils import get_column_letter
        from pathlib import Path
        import datetime

        if self._export_data is None:
            self._append_log("[WARN] No test result to export.")
            return

        # 多通道高低温测试：每通道一个 sheet 的专用导出
        if 'multi_ch_temp' in self._export_data:
            self._export_multi_ch_temp(self._export_data['multi_ch_temp'])
            return

        # FT Calibration Check：校准汇总 + 逐点数据的专用导出
        if 'ft_calib' in self._export_data:
            self._export_ft_calib(self._export_data['ft_calib'])
            return

        params   = self._export_data['params']
        raw      = self._export_data['raw']
        calib    = self._export_data['calibration']

        results_dir = Path(__file__).parent.parent.parent.parent / "Results"
        results_dir.mkdir(exist_ok=True)

        ts       = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = results_dir / f"GPADC_Result_{ts}.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "GPADC Result"

        hdr_fill   = PatternFill("solid", fgColor="0A1735")
        hdr_font   = Font(bold=True, color="00D39A", size=11)
        val_font   = Font(color="000000", size=10)
        sub_fill   = PatternFill("solid", fgColor="0D1F40")
        sub_font   = Font(bold=True, color="7E96BF", size=10)
        thin_side  = Side(style="thin", color="1B2847")
        thin_border = Border(left=thin_side, right=thin_side,
                             top=thin_side, bottom=thin_side)

        def _hdr(ws, row, col, text):
            c = ws.cell(row=row, column=col, value=text)
            c.fill = hdr_fill
            c.font = hdr_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
            return c

        def _val(ws, row, col, value):
            c = ws.cell(row=row, column=col, value=value)
            c.font = val_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
            if isinstance(value, float):
                c.number_format = "0.000000"
            return c

        def _sub(ws, row, col, text):
            c = ws.cell(row=row, column=col, value=text)
            c.fill = sub_fill
            c.font = sub_font
            c.alignment = Alignment(horizontal="left", vertical="center")
            c.border = thin_border
            return c

        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 18

        row = 1
        _hdr(ws, row, 1, "GPADC ADC Parameters")
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=4)
        row += 1

        param_items = [
            ("Linearity (R²)",   params.get("linearity",    0.0)),
            ("INL (LSB)",        params.get("inl",          0.0)),
            ("DNL (LSB)",        params.get("dnl",          0.0)),
            ("ENOB (bits)",      params.get("enob",         0.0)),
            ("Gain Error (%)",   params.get("gain_error",   0.0)),
            ("Offset Error (LSB)", params.get("offset_error", 0.0)),
        ]
        _sub(ws, row, 1, "Parameter")
        _sub(ws, row, 2, "Value")
        _sub(ws, row, 3, "Parameter")
        _sub(ws, row, 4, "Value")
        row += 1
        for i in range(0, len(param_items), 2):
            _val(ws, row, 1, param_items[i][0])
            _val(ws, row, 2, param_items[i][1])
            if i + 1 < len(param_items):
                _val(ws, row, 3, param_items[i + 1][0])
                _val(ws, row, 4, param_items[i + 1][1])
            row += 1

        row += 1

        chart_data_start_row = row
        _hdr(ws, row, 1, "Chart Data")
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=4)
        row += 1
        _sub(ws, row, 1, "Input Voltage (V)")
        _sub(ws, row, 2, "Mean Cali (V)")
        _sub(ws, row, 3, "Min Cali (V)")
        _sub(ws, row, 4, "Max Cali (V)")
        row += 1
        chart_data_body_start = row
        v_list   = calib['voltage']
        mc_list  = calib['mean_cali']
        mn_list  = calib['min_cali']
        mx_list  = calib['max_cali']
        for i in range(len(v_list)):
            _val(ws, row, 1, float(v_list[i]))
            _val(ws, row, 2, float(mc_list[i]))
            _val(ws, row, 3, float(mn_list[i]))
            _val(ws, row, 4, float(mx_list[i]))
            row += 1
        chart_data_end = row - 1

        chart_anchor_col = get_column_letter(6)
        if self._chart_image_bytes is not None:
            try:
                from openpyxl.drawing.image import Image as XLImage
                self._chart_image_bytes.seek(0)
                img = XLImage(self._chart_image_bytes)
                img.width  = 600
                img.height = 350
                ws.add_image(img, f"{chart_anchor_col}{chart_data_start_row}")
            except Exception as ex:
                self._append_log(f"[WARN] Embed chart image failed: {ex}")
                self._embed_native_chart(ws, chart_data_start_row, chart_data_body_start,
                                         chart_data_end, chart_anchor_col)
        else:
            self._embed_native_chart(ws, chart_data_start_row, chart_data_body_start,
                                     chart_data_end, chart_anchor_col)

        row += 2

        _hdr(ws, row, 1, "Raw Data")
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=4)
        row += 1
        _sub(ws, row, 1, "Input Voltage (V)")
        _sub(ws, row, 2, "ADC Mean (LSB)")
        _sub(ws, row, 3, "ADC Min (LSB)")
        _sub(ws, row, 4, "ADC Max (LSB)")
        row += 1
        rv_list  = raw['voltage']
        rm_list  = raw['mean']
        rn_list  = raw['min']
        rx_list  = raw['max']
        for i in range(len(rv_list)):
            _val(ws, row, 1, float(rv_list[i]))
            _val(ws, row, 2, float(rm_list[i]))
            _val(ws, row, 3, float(rn_list[i]))
            _val(ws, row, 4, float(rx_list[i]))
            row += 1

        row += 2

        _hdr(ws, row, 1, "Calibrated Data")
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=4)
        row += 1
        _sub(ws, row, 1, "Input Voltage (V)")
        _sub(ws, row, 2, "Mean Cali (V)")
        _sub(ws, row, 3, "Min Cali (V)")
        _sub(ws, row, 4, "Max Cali (V)")
        row += 1
        for i in range(len(v_list)):
            _val(ws, row, 1, float(v_list[i]))
            _val(ws, row, 2, float(mc_list[i]))
            _val(ws, row, 3, float(mn_list[i]))
            _val(ws, row, 4, float(mx_list[i]))
            row += 1

        wb.save(str(filename))
        self._append_log(f"[INFO] Result exported to: {filename}")

    def _embed_native_chart(self, ws, chart_data_start_row, chart_data_body_start,
                            chart_data_end, chart_anchor_col):
        from openpyxl.chart import LineChart, Reference
        chart = LineChart()
        chart.title        = "ADC Transfer Curve (Calibrated)"
        chart.style        = 10
        chart.y_axis.title = "Calibrated Voltage (V)"
        chart.x_axis.title = "Input Voltage (V)"
        chart.width        = 22
        chart.height       = 14
        ref_x    = Reference(ws, min_col=1, min_row=chart_data_body_start, max_row=chart_data_end)
        ref_mean = Reference(ws, min_col=2, min_row=chart_data_body_start - 1, max_row=chart_data_end)
        ref_min  = Reference(ws, min_col=3, min_row=chart_data_body_start - 1, max_row=chart_data_end)
        ref_max  = Reference(ws, min_col=4, min_row=chart_data_body_start - 1, max_row=chart_data_end)
        chart.add_data(ref_mean, titles_from_data=True)
        chart.add_data(ref_min,  titles_from_data=True)
        chart.add_data(ref_max,  titles_from_data=True)
        chart.set_categories(ref_x)
        chart.series[0].graphicalProperties.line.solidFill = "00D39A"
        chart.series[0].graphicalProperties.line.width     = 20000
        chart.series[1].graphicalProperties.line.solidFill = "F0A040"
        chart.series[2].graphicalProperties.line.solidFill = "F0A040"
        ws.add_chart(chart, f"{chart_anchor_col}{chart_data_start_row}")

    def _export_ft_calib(self, result):
        """FT Calibration Check 导出：Summary（K/B + 校准点确认 + 误差统计 + 判定）+ Sweep Data。"""
        import openpyxl
        from openpyxl.styles import (Font, PatternFill, Alignment,
                                     Border, Side)
        from openpyxl.utils import get_column_letter
        from pathlib import Path
        import datetime

        results_dir = Path(__file__).parent.parent.parent.parent / "Results"
        results_dir.mkdir(exist_ok=True)

        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = results_dir / f"GPADC_FT_CalibCheck_{ts}.xlsx"

        hdr_fill = PatternFill("solid", fgColor="0A1735")
        hdr_font = Font(bold=True, color="00D39A", size=11)
        val_font = Font(color="000000", size=10)
        sub_fill = PatternFill("solid", fgColor="0D1F40")
        sub_font = Font(bold=True, color="7E96BF", size=10)
        thin_side = Side(style="thin", color="1B2847")
        thin_border = Border(left=thin_side, right=thin_side,
                             top=thin_side, bottom=thin_side)

        def _hdr(ws, row, col, text):
            c = ws.cell(row=row, column=col, value=text)
            c.fill = hdr_fill
            c.font = hdr_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
            return c

        def _val(ws, row, col, value):
            c = ws.cell(row=row, column=col, value=value)
            c.font = val_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
            if isinstance(value, float):
                c.number_format = "0.000000"
            return c

        def _sub(ws, row, col, text):
            c = ws.cell(row=row, column=col, value=text)
            c.fill = sub_fill
            c.font = sub_font
            c.alignment = Alignment(horizontal="left", vertical="center")
            c.border = thin_border
            return c

        mode = result.get('mode', '')
        stats = result.get('stats') or {}
        cons_stats = result.get('cons_stats')
        (v1, c1), (v2, c2) = result.get('calib_points') or [(0.0, 0.0), (0.0, 0.0)]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 20

        row = 1
        _hdr(ws, row, 1, "FT Calibration Check Summary")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1
        summary_items = [
            ("Mode", mode),
            ("K (LSB/V)", result.get('k')),
            ("B (LSB)", result.get('b')),
            ("Calib Point 1 (V)", v1),
            ("Calib Raw 1 (LSB)", c1),
            ("Calib Point 2 (V)", v2),
            ("Calib Raw 2 (LSB)", c2),
            ("Raw Tol (±LSB)", result.get('raw_tol_lsb')),
            ("Error Limit (±mV)", result.get('err_limit_mv')),
            ("Divider Ratio (×)", result.get('divider_ratio')),
        ]
        for name, value in summary_items:
            _sub(ws, row, 1, name)
            _val(ws, row, 2, value)
            row += 1

        point_checks = result.get('point_checks') or []
        if point_checks:
            row += 1
            _hdr(ws, row, 1, "Calib Point Check (IIC)")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            row += 1
            for col, name in enumerate(
                    ("Voltage (V)", "Set Voltage (V)", "Calib Code (LSB)",
                     "Measured (LSB)", "Dev (LSB)", "Result"), 1):
                _sub(ws, row, col, name)
            row += 1
            for pc in point_checks:
                _val(ws, row, 1, pc.get('voltage'))
                _val(ws, row, 2, pc.get('set_voltage'))
                _val(ws, row, 3, pc.get('calib_code'))
                _val(ws, row, 4, pc.get('measured'))
                _val(ws, row, 5, pc.get('dev_lsb'))
                _val(ws, row, 6, "PASS" if pc.get('passed') else "FAIL")
                row += 1

        row += 1
        _hdr(ws, row, 1, "Error Statistics")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1
        stat_items = [
            ("Max Abs Error (mV)", stats.get('max_abs_mv')),
            ("Avg Error (mV)", stats.get('avg_mv')),
            ("RMS Error (mV)", stats.get('rms_mv')),
            ("Sweep Points", stats.get('count')),
        ]
        if cons_stats is not None:
            stat_items.extend([
                ("Cons Max Abs (mV)", cons_stats.get('max_abs_mv')),
                ("Cons Avg (mV)", cons_stats.get('avg_mv')),
                ("Cons RMS (mV)", cons_stats.get('rms_mv')),
            ])
        stat_items.append(("Verdict", "PASS" if result.get('passed') else "FAIL"))
        for name, value in stat_items:
            _sub(ws, row, 1, name)
            _val(ws, row, 2, value)
            row += 1

        ws2 = wb.create_sheet("Sweep Data")
        is_uart = mode == 'UART'
        headers = ["Set Voltage (V)", "DUT Voltage (V)", "Raw Mean (LSB)",
                   "Calib Voltage (V)", "Error (mV)"]
        if is_uart:
            headers += ["DUT Volt (mV)", "Consistency (mV)"]
        for col, name in enumerate(headers, 1):
            _sub(ws2, row=1, col=col, text=name)
            ws2.column_dimensions[get_column_letter(col)].width = 20
        voltage = result.get('voltage') or []
        actual_voltage = result.get('actual_voltage') or voltage
        raw_mean = result.get('raw_mean') or []
        cal_volt = result.get('cal_volt') or []
        err_mv = result.get('err_mv') or []
        dut_volt = result.get('dut_volt_mv') or []
        cons_mv = result.get('cons_mv') or []
        for i in range(len(voltage)):
            r = i + 2
            _val(ws2, r, 1, float(voltage[i]))
            _val(ws2, r, 2, float(actual_voltage[i]))
            _val(ws2, r, 3, float(raw_mean[i]))
            _val(ws2, r, 4, float(cal_volt[i]))
            _val(ws2, r, 5, float(err_mv[i]))
            if is_uart:
                _val(ws2, r, 6, float(dut_volt[i]))
                _val(ws2, r, 7, float(cons_mv[i]))

        wb.save(filename)
        self._append_log(f"[INFO] FT Calibration Check 结果已导出: {filename}")

    def _export_multi_ch_temp(self, result):
        """多通道高低温测试导出：Summary 汇总 + 每通道一个 sheet（原始 + 标定数据）。"""
        import openpyxl
        from openpyxl.styles import (Font, PatternFill, Alignment,
                                     Border, Side)
        from openpyxl.utils import get_column_letter
        from pathlib import Path
        import datetime
        import re

        results_dir = Path(__file__).parent.parent.parent.parent / "Results"
        results_dir.mkdir(exist_ok=True)

        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = results_dir / f"GPADC_MultiChTemp_{ts}.xlsx"

        hdr_fill = PatternFill("solid", fgColor="0A1735")
        hdr_font = Font(bold=True, color="00D39A", size=11)
        val_font = Font(color="000000", size=10)
        sub_fill = PatternFill("solid", fgColor="0D1F40")
        sub_font = Font(bold=True, color="7E96BF", size=10)
        thin_side = Side(style="thin", color="1B2847")
        thin_border = Border(left=thin_side, right=thin_side,
                             top=thin_side, bottom=thin_side)

        def _hdr(ws, row, col, text):
            c = ws.cell(row=row, column=col, value=text)
            c.fill = hdr_fill
            c.font = hdr_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
            return c

        def _val(ws, row, col, value):
            c = ws.cell(row=row, column=col, value=value)
            c.font = val_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
            if isinstance(value, float):
                c.number_format = "0.000000"
            return c

        def _sub(ws, row, col, text):
            c = ws.cell(row=row, column=col, value=text)
            c.fill = sub_fill
            c.font = sub_font
            c.alignment = Alignment(horizontal="left", vertical="center")
            c.border = thin_border
            return c

        temp_list = result.get('temp') or []
        channels = result.get('channels') or []

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 22

        row = 1
        _hdr(ws, row, 1, "GPADC Multi-Ch Temp Test Summary")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
        meta_items = [
            ("Test Time", ts),
            ("Temp Range (°C)", f"{temp_list[0]:.1f} ~ {temp_list[-1]:.1f}" if temp_list else "N/A"),
            ("Temp Points", len(temp_list)),
            ("Sample Count", result.get('sample_cnt', 0)),
            ("Channel Count", len(channels)),
        ]
        for key, value in meta_items:
            _sub(ws, row, 1, key)
            _val(ws, row, 2, value)
            row += 1

        row += 1
        _hdr(ws, row, 1, "Channels")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
        for col, text in enumerate(("Name", "Sweep", "Voltage Range (V)", "Data Rows"), start=1):
            _sub(ws, row, col, text)
        row += 1
        for ch in channels:
            _val(ws, row, 1, ch.get('name') or "CH")
            _val(ws, row, 2, "Yes" if ch.get('sweep') else "No")
            if ch.get('sweep'):
                vpts = ch.get('voltage') or []
                vrange = f"{vpts[0]:.3f} ~ {vpts[-1]:.3f}" if vpts else "N/A"
            else:
                vrange = "(fixed-point)"
            _val(ws, row, 3, vrange)
            _val(ws, row, 4, len(ch.get('mean') or []))
            row += 1

        if self._chart_image_bytes is not None:
            try:
                from openpyxl.drawing.image import Image as XLImage
                self._chart_image_bytes.seek(0)
                img = XLImage(self._chart_image_bytes)
                img.width = 720
                img.height = 420
                ws.add_image(img, f"{get_column_letter(6)}1")
            except Exception as ex:
                self._append_log(f"[WARN] Embed chart image failed: {ex}")

        for ch in channels:
            ch_name = ch.get('name') or "CH"
            sheet_name = re.sub(r"[\[\]:*?/\\]", "_", f"CH_{ch_name}")[:31]
            ws = wb.create_sheet(sheet_name)
            for col_idx in range(1, 12):
                ws.column_dimensions[get_column_letter(col_idx)].width = 14

            row = 1
            _hdr(ws, row, 1, f"Channel: {ch_name} ({'voltage sweep' if ch.get('sweep') else 'fixed-point'})")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            row += 2

            if ch.get('sweep'):
                vpts = ch.get('voltage') or []
                matrices = [("Mean (LSB)", ch.get('mean') or []),
                            ("Min (LSB)", ch.get('min') or []),
                            ("Max (LSB)", ch.get('max') or [])]
                if ch.get('mean_cali'):
                    matrices.append(("Mean Cali (V)", ch['mean_cali']))
                for tag, matrix in matrices:
                    _hdr(ws, row, 1, tag)
                    ws.merge_cells(start_row=row, start_column=1,
                                   end_row=row, end_column=max(2, len(vpts) + 1))
                    row += 1
                    _sub(ws, row, 1, "Temp \\ Voltage (V)")
                    for j, v in enumerate(vpts):
                        _sub(ws, row, 2 + j, f"{v:.3f}")
                    row += 1
                    for i, t in enumerate(temp_list):
                        if i >= len(matrix):
                            break
                        _val(ws, row, 1, f"{t:.1f}")
                        for j in range(len(matrix[i])):
                            _val(ws, row, 2 + j, float(matrix[i][j]))
                        row += 1
                    row += 1
            else:
                _sub(ws, row, 1, "Temp (°C)")
                _sub(ws, row, 2, "ADC Mean (LSB)")
                _sub(ws, row, 3, "ADC Min (LSB)")
                _sub(ws, row, 4, "ADC Max (LSB)")
                row += 1
                mean_list = ch.get('mean') or []
                min_list = ch.get('min') or []
                max_list = ch.get('max') or []
                for i in range(len(mean_list)):
                    t = temp_list[i] if i < len(temp_list) else i
                    _val(ws, row, 1, float(t))
                    _val(ws, row, 2, float(mean_list[i]))
                    if i < len(min_list):
                        _val(ws, row, 3, float(min_list[i]))
                    if i < len(max_list):
                        _val(ws, row, 4, float(max_list[i]))
                    row += 1

        wb.save(str(filename))
        self._append_log(f"[INFO] Result exported to: {filename}")
